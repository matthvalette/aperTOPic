"""Label sets known from the labelled data in files.openparldata.ch/hackapertus.

Implements `TaxonomyProvider`. Topic names come from the data; the scope descriptions of the
federal subject areas are written by hand (see SUBJECT_AREA_SCOPES) because the descriptors
in the data predate the 2012 scheme change and describe the old 20-category scheme.
"""

from __future__ import annotations

import csv
import re
import unicodedata
from pathlib import Path

from ..domain.models import Taxonomy, TaxonomyKind, Topic

CURIA = "curia_subject_areas"
SG = "sg_law_collection"
TG = "tg_subject_codes"
DIGITALISATION = "digitalisation"
APS = "aps_topics"

# English name and scope of each Curia Vista subject area (Sachbereich), current scheme.
# Hand-written glosses for the classifier prompt, not official texts. Edit freely.
SUBJECT_AREA_SCOPES: dict[str, tuple[str, str]] = {
    "Beschäftigung und Arbeit": ("Employment and labour", "Labour market, jobs and unemployment, wages, working conditions and working time, labour law, protection of workers, wage dumping and accompanying measures to free movement."),
    "Bildung": ("Education", "Schools, vocational education and apprenticeships, universities and universities of applied sciences, further education, recognition of diplomas, education exchange programmes."),
    "Energie": ("Energy", "Electricity supply and grids, nuclear power and radioactive waste, renewable energy, energy efficiency and savings, energy strategy, oil and gas."),
    "Europapolitik": ("European policy", "Relations between Switzerland and the EU, bilateral agreements and negotiations, institutional questions, Schengen/Dublin, EU programmes, Council of Europe."),
    "Finanzwesen": ("Public finance and financial markets", "Federal budget, spending and accounts, debt brake, financial equalisation between cantons, banks, the financial centre, financial market supervision, the National Bank. Taxes are a separate subject area."),
    "Gerichtswesen": ("Judiciary", "Federal courts, judges and their election, organisation and procedure of courts, access to justice and legal aid."),
    "Gesundheit": ("Health", "Health insurance and premiums, hospitals, doctors and nursing staff, medicines, prevention, epidemics and pandemics, addiction, patients' rights, health costs."),
    "Internationale Politik": ("International relations", "Foreign policy, diplomacy, relations with other countries, development cooperation, humanitarian aid, peace policy and good offices, international organisations, sanctions."),
    "Internationales Recht": ("International law", "International treaties and conventions, international legal assistance, return of illicit assets of politically exposed persons, international courts, international humanitarian law."),
    "Kultur": ("Culture", "Promotion of culture, languages and multilingualism, heritage and monuments, film, arts, museums, religious communities."),
    "Landwirtschaft": ("Agriculture", "Agricultural policy and direct payments, farming and livestock, agricultural products and trade, plant protection products, forestry, food production."),
    "Medien und Kommunikation": ("Media and communication", "Press, radio and television including the SRG SSR and licence fees, telecommunications and mobile networks, broadband, postal services, the internet and information technology."),
    "Menschenrechte": ("Human rights", "Fundamental and human rights, discrimination (racism, antisemitism), equal treatment, freedom of religion, protection of minorities, human rights abroad."),
    "Migration": ("Migration", "Asylum and refugees, immigration and residence of foreigners, deportation, naturalisation, integration, control of immigration."),
    "Parlament": ("Parliament", "Parliamentary procedure, members of parliament and their pay and infrastructure, committees and delegations, parliamentary services, the Federal Palace."),
    "Raumplanung und Wohnungswesen": ("Spatial planning and housing", "Land-use planning, building zones, second homes, housing construction, rents and tenancy, home ownership, mortgages, land purchases by foreigners."),
    "Recht Allgemein": ("Law in general", "Legal questions that do not belong to civil law, criminal law, international law or the judiciary, for example legislation in general, legal certainty, liability of the state."),
    "Sicherheitspolitik": ("Security policy", "Army and armed forces, military procurement such as fighter jets, intelligence service, civil protection, terrorism, police cooperation, public safety, arms exports, civilian service."),
    "Soziale Fragen": ("Social questions", "Family policy and childcare, equality between women and men, old age, disability, poverty and social assistance, youth, social cohesion."),
    "Sozialer Schutz": ("Social security", "Social insurances: old-age and survivors' insurance (AHV/AVS), occupational pensions and pension funds, disability insurance, unemployment insurance, supplementary benefits, family allowances."),
    "Staatspolitik": ("Political institutions", "Federal administration and its staff, the Federal Council, federalism and the cantons, elections and votes, popular rights, official languages, digital administration, public procurement."),
    "Steuer": ("Taxation", "Federal, cantonal and municipal taxes: income, wealth and corporate tax, VAT, withholding tax, tax treaties and double taxation, tax evasion, lump-sum taxation."),
    "Strafrecht": ("Criminal law", "Criminal code and criminal procedure, offences and offenders, violence and sexual offences, sentences and prisons, crime prevention, youth crime."),
    "Umwelt": ("Environment", "Climate policy and CO2, protection of water, soil and air, biodiversity and nature conservation, animal protection and wildlife, hunting, noise, waste, natural hazards."),
    "Verkehr": ("Transport", "Rail and the SBB, roads and motorways, road traffic rules and road safety, public transport, freight transport and transalpine traffic, aviation and airports, cycling."),
    "Wirtschaft": ("Economy", "Economic policy, competition and cartels, companies and SMEs, foreign trade and free trade agreements, consumer protection, prices, tourism, economic location."),
    "Wissenschaft und Forschung": ("Science and research", "Research policy and research funding, the National Science Foundation, research at universities, EU research programmes, innovation."),
    "Zivilrecht": ("Civil law", "Family and inheritance law, contract law, company law, property law, private insurance law, adult and child protection, debt enforcement."),
}

DIGITALISATION_SCOPE = (
    "Digitalisation in the broad editorial sense of PolitMonitor: e-government and digital "
    "public services, IT projects of the administration, cybersecurity, artificial intelligence, "
    "data protection and data use, digital media and platforms, broadband and telecommunications "
    "infrastructure, the digital economy."
)


def slugify(value: str) -> str:
    ascii_value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", ascii_value.lower()).strip("-")


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def build_curia(data_dir: Path) -> Taxonomy:
    topics = []
    for row in _read_csv(data_dir / "che_topics_vocabulary.csv"):
        english, scope = SUBJECT_AREA_SCOPES.get(row["topic_de"], (row["topic_de"], ""))
        topics.append(
            Topic(
                id=row["topic_id"],
                name=row["topic_de"],
                names={"de": row["topic_de"], "fr": row["topic_fr"], "it": row["topic_it"], "en": english},
                description=scope or None,
            )
        )
    return Taxonomy(
        id=CURIA,
        name="Curia Vista subject areas",
        kind=TaxonomyKind.MULTI_LABEL,
        description="The 28 subject areas the federal Parliamentary Services assign to every federal affair, usually one to three per affair.",
        source="che_topics_vocabulary.csv",
        topics=topics,
        # Best micro-F1 on the dev half, jev-1.14.0, 300 affairs 2018-2026 (evaluation_runs/).
        default_threshold=0.75,
        native_bodies=["CHE"],
    )


# A few Thurgau rows join two codes in one label ("..., 61.1 Finanzhaushalt ..."). They are
# normalisation artefacts with one affair each, so they are left out of the scheme.
_SECOND_CODE = re.compile(r",\s*\d{1,2}\.\d+\s")


def build_cantonal(data_dir: Path, body_key: str, taxonomy_id: str, name: str, description: str) -> Taxonomy:
    topics = []
    for row in _read_csv(data_dir / "cantonal_keywords_vocabulary.csv"):
        if row["body_key"] != body_key:
            continue
        code = row.get("keyword_code") or None
        label = row["keyword_de"]
        if _SECOND_CODE.search(label):
            continue
        topics.append(
            Topic(
                id=code or slugify(label),
                name=f"{code} {label}" if code else label,
                names={"de": label},
                code=code,
            )
        )
    topics.sort(key=lambda t: t.id)
    return Taxonomy(
        id=taxonomy_id,
        name=name,
        kind=TaxonomyKind.SINGLE_LABEL,
        description=description,
        source="cantonal_keywords_vocabulary.csv",
        topics=topics,
        native_bodies=[body_key],
    )


def build_aps(data_dir: Path) -> Taxonomy:
    """Année Politique Suisse: domain > chapter > topic. No affair carries these labels."""
    from openpyxl import load_workbook

    workbook = load_workbook(data_dir / "aps_topics.xlsx", read_only=True)
    rows = list(workbook["Topics"].iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    topics = []
    for values in rows[1:]:
        row = dict(zip(header, values))
        if not row.get("topic_de"):
            continue
        topics.append(
            Topic(
                id=str(row["code"]),
                name=str(row["topic_de"]),
                names={"de": str(row["topic_de"]), "fr": str(row.get("topic_fr") or "")},
                code=str(row["code"]),
                group=str(row["chapter_de"]),
                description=f"APS chapter {row['chapter_de']} ({row.get('chapter_fr')}), "
                f"domain {row['domain_de']} ({row.get('domain_fr')}).",
            )
        )
    return Taxonomy(
        id=APS,
        name="Année Politique Suisse topics",
        kind=TaxonomyKind.MULTI_LABEL,
        description="The challenge's target scheme: 165 topics in 22 chapters. No labelled affairs exist, "
        "so its accuracy cannot be measured on this data. Chapter scores are rolled up from topics.",
        source="aps_topics.xlsx",
        topics=topics,
        # Not tuned (no labels): borrowed from the subject areas, which use the same question.
        default_threshold=0.75,
    )


def build_digitalisation() -> Taxonomy:
    return Taxonomy(
        id=DIGITALISATION,
        name="Digitalisation (ad-hoc topic)",
        kind=TaxonomyKind.BINARY,
        description="The PolitMonitor editorial tag behind the DVS list of political affairs on digitalisation.",
        source="politmonitor_digitalisierung.csv",
        topics=[
            Topic(
                id="digitalisation",
                name="Digitalisierung",
                names={"de": "Digitalisierung", "fr": "Numérisation", "it": "Digitalizzazione", "en": "Digitalisation"},
                description=DIGITALISATION_SCOPE,
            )
        ],
    )


class LabelledDataTaxonomies:
    """Builds every label set once from the CSV vocabularies."""

    def __init__(self, data_dir: Path) -> None:
        if not (data_dir / "che_topics_vocabulary.csv").exists():
            raise FileNotFoundError(f"Labelled data not found in {data_dir}. Set DATA_DIR.")
        taxonomies = [
            build_curia(data_dir),
            build_cantonal(
                data_dir, "SG", SG, "St. Gallen law-collection areas",
                "The 10 sections of the St. Gallen systematic law collection. Each affair gets exactly one.",
            ),
            build_cantonal(
                data_dir, "TG", TG, "Thurgau subject codes",
                "Thurgau's systematic subject register. Each affair gets exactly one code.",
            ),
            build_aps(data_dir),
            build_digitalisation(),
        ]
        self._by_id = {t.id: t for t in taxonomies}

    def list(self) -> list[Taxonomy]:
        return list(self._by_id.values())

    def get(self, taxonomy_id: str) -> Taxonomy:
        try:
            return self._by_id[taxonomy_id]
        except KeyError:
            raise KeyError(f"Unknown taxonomy {taxonomy_id!r}. Known: {', '.join(self._by_id)}") from None


def custom_taxonomy(topics: list[tuple[str, str | None]]) -> Taxonomy:
    """An ad-hoc taxonomy from user-named topics. No training data needed."""
    seen: set[str] = set()
    built = []
    for name, description in topics:
        topic_id = slugify(name) or f"topic-{len(built) + 1}"
        while topic_id in seen:
            topic_id += "-x"
        seen.add(topic_id)
        built.append(Topic(id=topic_id, name=name, description=description))
    return Taxonomy(
        id="custom",
        name="Ad-hoc topics",
        kind=TaxonomyKind.MULTI_LABEL,
        description="Topics named in the request.",
        source="request",
        topics=built,
    )
